"""Job application tracker: suivi-candidatures.xlsx at the project root.

Usage:
    python tracker.py upsert --company "Acme" --position "Stage Développeur IA" \
        --domain "IA agentique / EdTech" --url https://... [--status Prepared] [--applied 2026-10-14] \
        [--location Paris] [--contract "Stage 4-6 mois"] [--fit 72] [--folder output/...] [--notes "..."]
    python tracker.py status --url https://... --status Applied [--applied 2026-10-14] [--notes "..."]
    python tracker.py list

Rows are matched by URL (or by company + position when there's no URL). Re-running upsert updates the row.
Statuses: To apply, Prepared, Applied, Interview, Offer, Rejected, Withdrawn.
"""
from __future__ import annotations

import argparse
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[4]
XLSX = ROOT / "suivi-candidatures.xlsx"
SHEET = "Candidatures"
COLUMNS = [  # (header, key, width)
    ("#", "id", 5), ("Company", "company", 24), ("Position", "position", 38), ("Domain", "domain", 26),
    ("Location", "location", 12), ("Contract", "contract", 16), ("Link", "url", 14), ("Status", "status", 12),
    ("Date added", "added", 12), ("Date applied", "applied", 13), ("Fit", "fit", 7), ("Folder", "folder", 34),
    ("Notes / next step", "notes", 50),
]
STATUSES = ["To apply", "Prepared", "Applied", "Interview", "Offer", "Rejected", "Withdrawn"]
STATUS_COLORS = {"To apply": "EDEDED", "Prepared": "DCE6F1", "Applied": "FFF2CC", "Interview": "E2EFDA",
                 "Offer": "A9D08E", "Rejected": "F8CBAD", "Withdrawn": "D9D9D9"}
ACCENT = "1F3A5F"
COL = {key: i + 1 for i, (_, key, _) in enumerate(COLUMNS)}


def open_book():
    if XLSX.exists():
        wb = load_workbook(XLSX)
        return wb, wb[SHEET]
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([h for h, _, _ in COLUMNS])
    return wb, ws


def parse_date(value: str | None):
    if not value:
        return None
    if value.lower() == "today":
        return date.today()
    return datetime.strptime(value, "%Y-%m-%d").date()


def find_row(ws, url: str | None, company: str | None, position: str | None) -> int | None:
    for r in range(2, ws.max_row + 1):
        cell_url = ws.cell(r, COL["url"]).hyperlink.target if ws.cell(r, COL["url"]).hyperlink else None
        if url and cell_url == url:
            return r
        if not url and company and position and ws.cell(r, COL["company"]).value == company \
                and ws.cell(r, COL["position"]).value == position:
            return r
    return None


def style(ws):
    thin = Side(style="thin", color="D0D0D0")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        c = ws.cell(1, i)
        c.font, c.fill = header_font, PatternFill("solid", fgColor=ACCENT)
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 22
    for r in range(2, ws.max_row + 1):
        for i in range(1, len(COLUMNS) + 1):
            c = ws.cell(r, i)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical="top", wrap_text=i in (COL["company"], COL["position"], COL["domain"], COL["notes"], COL["folder"]))
            if c.font is None or not c.hyperlink:
                c.font = Font(name="Arial", size=10)
        for key in ("added", "applied"):
            ws.cell(r, COL[key]).number_format = "yyyy-mm-dd"
        link = ws.cell(r, COL["url"])
        if link.hyperlink:
            link.font = Font(name="Arial", size=10, color="0563C1", underline="single")
    ws.freeze_panes = "C2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    last = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last}{max(ws.max_row, 2)}"

    # Status drop-down + colour per status (rebuilt each time).
    ws.data_validations.dataValidation = []
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    status_col = get_column_letter(COL["status"])
    dv.add(f"{status_col}2:{status_col}500")
    ws.add_data_validation(dv)
    ws.conditional_formatting = type(ws.conditional_formatting)()
    for status, color in STATUS_COLORS.items():
        ws.conditional_formatting.add(f"A2:{last}500", FormulaRule(
            formula=[f'${status_col}2="{status}"'], fill=PatternFill("solid", fgColor=color)))


def save(wb, ws):
    style(ws)
    try:
        wb.save(XLSX)
    except PermissionError:
        sys.exit(f"Cannot write {XLSX.name}: close it in Excel and run again.")


def upsert(a):
    wb, ws = open_book()
    r = find_row(ws, a.url, a.company, a.position)
    if r is None:
        r = ws.max_row + 1
        ws.cell(r, COL["id"], r - 1)
        ws.cell(r, COL["added"], date.today())
        ws.cell(r, COL["status"], a.status or "To apply")
    values = {"company": a.company, "position": a.position, "domain": a.domain, "location": a.location,
              "contract": a.contract, "status": a.status, "applied": parse_date(a.applied), "fit": a.fit,
              "folder": a.folder, "notes": a.notes}
    for key, v in values.items():
        if v is not None:
            ws.cell(r, COL[key], v)
    if a.url:
        c = ws.cell(r, COL["url"], "Open offer")
        c.hyperlink = a.url
    if (a.status == "Applied") and not ws.cell(r, COL["applied"]).value:
        ws.cell(r, COL["applied"], date.today())
    save(wb, ws)
    print(f"Row {r - 1} saved: {ws.cell(r, COL['company']).value} | {ws.cell(r, COL['status']).value} -> {XLSX}")


def set_status(a):
    wb, ws = open_book()
    r = find_row(ws, a.url, a.company, a.position)
    if r is None:
        sys.exit("No matching row. Use upsert first.")
    ws.cell(r, COL["status"], a.status)
    if a.applied or (a.status == "Applied" and not ws.cell(r, COL["applied"]).value):
        ws.cell(r, COL["applied"], parse_date(a.applied) or date.today())
    if a.notes:
        ws.cell(r, COL["notes"], a.notes)
    save(wb, ws)
    print(f"Row {r - 1}: {ws.cell(r, COL['company']).value} -> {a.status}")


def list_rows(_):
    if not XLSX.exists():
        print("No tracker yet.")
        return
    ws = load_workbook(XLSX)[SHEET]
    for r in range(2, ws.max_row + 1):
        v = {key: ws.cell(r, COL[key]).value for _, key, _ in COLUMNS}
        applied = v["applied"].strftime("%Y-%m-%d") if v["applied"] else "-"
        print(f"{v['id']:>2} | {v['status']:<9} | {applied} | {v['company']} — {v['position']} | {v['domain']}")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    u = sub.add_parser("upsert")
    s = sub.add_parser("status")
    for p in (u, s):
        p.add_argument("--url")
        p.add_argument("--company")
        p.add_argument("--position")
        p.add_argument("--applied", help="YYYY-MM-DD or 'today'")
        p.add_argument("--notes")
    u.add_argument("--domain")
    u.add_argument("--location")
    u.add_argument("--contract")
    u.add_argument("--status", choices=STATUSES)
    u.add_argument("--fit", type=int)
    u.add_argument("--folder")
    s.add_argument("--status", choices=STATUSES, required=True)
    sub.add_parser("list")
    a = ap.parse_args()
    {"upsert": upsert, "status": set_status, "list": list_rows}[a.cmd](a)


if __name__ == "__main__":
    main()
