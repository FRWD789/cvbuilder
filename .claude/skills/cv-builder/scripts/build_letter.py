"""Render a cover letter (lettre-de-motivation.md) to DOCX + PDF, in the same style as the CV (ats-modern).

Usage:
    python build_letter.py <folder>/lettre-de-motivation.md --company "Acme" \
        --subject "Candidature — Stage Développeur Full Stack (réf. 1234)" [--city Paris] [--name FILE_STEM] [--lang fr|en]

The letter body is the text between the first two lines made of "---" in the .md file.
Sender details (name, email, phone, links) come from profile/master-profile.json.
Writes <FILE_STEM>.docx and .pdf next to the .md (default stem: <NAME>_Lettre_<Company>, or <NAME>_Cover_Letter_<Company> in English).
--lang en switches the fixed labels (recipient line, date, "Subject:", closing) to English.
"""
from __future__ import annotations

import argparse
import json
import re
import unicodedata
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from build_cv import count_pdf_pages, to_pdf  # noqa: E402

ROOT = Path(__file__).resolve().parents[4]
PROFILE = ROOT / "profile" / "master-profile.json"
if not PROFILE.exists():  # fresh clone: fall back to the example profile
    PROFILE = ROOT / "profile" / "master-profile.example.json"
MONTHS_EN = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
MONTHS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
TEXT, ACCENT, FONT = RGBColor(0x22, 0x22, 0x22), RGBColor(0x1F, 0x3A, 0x5F), "Arial"
INLINE = re.compile(r"\*\*(.+?)\*\*")


def letter_body(md: str) -> list[str]:
    parts = re.split(r"^-{3,}\s*$", md, flags=re.M)
    if len(parts) < 3:
        sys.exit("No letter found: the body must sit between two '---' lines.")
    blocks = [b.strip() for b in re.split(r"\n\s*\n", parts[1].strip()) if b.strip()]
    # Drop the closing + signature block(s); the template adds its own.
    out = []
    for b in blocks:
        if re.match(r"(?i)^(bien |très )?cordialement|^sincères salutations|^(kind|best|warm) regards|^(yours )?sincerely", b):
            break
        out.append(" ".join(line.strip() for line in b.splitlines()))
    return out


def para(doc, text="", size=10.5, bold=False, color=TEXT, align=None, before=0, after=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.15
    if align is not None:
        p.alignment = align
    pos = 0
    for m in INLINE.finditer(text):  # support **bold**
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()]); r.font.size, r.font.color.rgb, r.bold = Pt(size), color, bold or None
        r = p.add_run(m.group(1)); r.font.size, r.font.color.rgb, r.bold = Pt(size), color, True
        pos = m.end()
    if pos < len(text) or not text:
        r = p.add_run(text[pos:]); r.font.size, r.font.color.rgb, r.bold = Pt(size), color, bold or None
    return p


def build(md_path: Path, company: str, subject: str, city: str, stem: str | None, lang: str = "fr"):
    en = lang == "en"
    b = json.loads(PROFILE.read_text(encoding="utf-8"))["basics"]
    city = city or (b.get("location") or "").split(",")[0].strip() or "Paris"
    body = letter_body(md_path.read_text(encoding="utf-8"))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name, st.font.size, st.font.color.rgb = FONT, Pt(10.5), TEXT
    rf = st.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), FONT)

    # Sender block (left), like the CV header
    para(doc, b["name"], size=16, bold=True, color=ACCENT)
    contact = " | ".join(x for x in (b.get("email"), b.get("phone"), b.get("location")) if x)
    para(doc, contact, size=10)
    links = " | ".join(l["label"] for l in b.get("links", []))
    if links:
        para(doc, links, size=10)

    # Recipient + date (right-aligned)
    para(doc, company, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, before=18)
    para(doc, "Recruitment Team" if en else "Service Recrutement", align=WD_ALIGN_PARAGRAPH.RIGHT)
    today = date.today()
    dated = f"{city}, {today.day} {MONTHS_EN[today.month - 1]} {today.year}" if en else f"{city}, le {today.day} {MONTHS_FR[today.month - 1]} {today.year}"
    para(doc, dated, align=WD_ALIGN_PARAGRAPH.RIGHT, before=10)

    # Subject
    para(doc, f"**Subject:** {subject}" if en else f"**Objet :** {subject}", before=18, after=12)

    for i, block in enumerate(body):
        para(doc, block, after=8 if i else 10)

    para(doc, "Kind regards," if en else "Bien cordialement,", before=6)
    para(doc, b["name"], bold=True, before=18)

    props = doc.core_properties
    props.title, props.author = f"{'Cover letter' if en else 'Lettre de motivation'} — {company}", b["name"]

    ascii_company = unicodedata.normalize("NFKD", company).encode("ascii", "ignore").decode()
    stem = stem or f"{b['name'].upper().replace(' ', '_')}_{'Cover_Letter' if en else 'Lettre'}_{re.sub(r'[^A-Za-z0-9]+', '_', ascii_company).strip('_')}"
    docx_path = md_path.parent / f"{stem}.docx"
    doc.save(docx_path)
    print(f"DOCX: {docx_path}")
    pdf = to_pdf(docx_path)
    if pdf:
        pages = count_pdf_pages(pdf)
        print(f"PDF:  {pdf}  ({pages} page{'s' if pages != 1 else ''})")
        if pages > 1:
            print("WARNING: the letter is longer than one page: shorten it.", file=sys.stderr)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("md", type=Path)
    ap.add_argument("--company", required=True)
    ap.add_argument("--subject", required=True)
    ap.add_argument("--city", help="default: the city in basics.location")
    ap.add_argument("--name", help="output file stem")
    ap.add_argument("--lang", choices=["fr", "en"], default="fr")
    a = ap.parse_args()
    build(a.md, a.company, a.subject, a.city, a.name, a.lang)


if __name__ == "__main__":
    main()
